from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\laragon\www\HRIS")
OUTPUT = ROOT / "artifacts" / "STRH_HRIS_Technical_and_Implementation_Proposal_Draft.docx"
LOGO = ROOT / "src" / "assets" / "branding" / "STRH-logo.png"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
STRH_BLUE = "0033A0"
GOLD = "F4B000"
INK = "1F2937"
MUTED = "667085"
LIGHT = "F4F6F9"
LIGHT_BLUE = "EEF4FB"
PALE_GOLD = "FFF6D8"
PALE_RED = "FDECEC"
WHITE = "FFFFFF"
BLACK = "000000"
GREEN = "246B45"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table, color="CBD5E1", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths, indent=120):
    if sum(widths) != 9360:
        raise ValueError(f"Table widths must total 9360 DXA, got {sum(widths)}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        prevent_row_split(row)
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Calibri", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_border(paragraph, color=GOLD, size=12):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    return run


def setup_numbering(doc):
    numbering = doc.part.numbering_part.element

    def add_num(abstract_id, num_id, fmt, text, left=540, hanging=280):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(left))
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), str(hanging))
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "290")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abs_id = OxmlElement("w:abstractNumId")
        abs_id.set(qn("w:val"), str(abstract_id))
        num.append(abs_id)
        numbering.append(num)

    add_num(90, 90, "bullet", "•")
    add_num(91, 91, "decimal", "%1.")


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.333
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    settings = {
        "Title": (28, BLACK, 0, 6),
        "Subtitle": (14, MUTED, 0, 8),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in settings.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for style_name in ("Proposal Kicker", "Table Text", "Table Header", "Callout", "Small Note"):
        if style_name not in styles:
            styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)

    kicker = styles["Proposal Kicker"]
    kicker.font.name = "Calibri"
    kicker.font.size = Pt(10)
    kicker.font.bold = True
    kicker.font.color.rgb = RGBColor.from_string(STRH_BLUE)
    kicker.paragraph_format.space_after = Pt(4)
    kicker.paragraph_format.keep_with_next = True

    table_text = styles["Table Text"]
    table_text.font.name = "Calibri"
    table_text.font.size = Pt(9.5)
    table_text.font.color.rgb = RGBColor.from_string(INK)
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(2)
    table_text.paragraph_format.line_spacing = 1.1

    table_header = styles["Table Header"]
    table_header.font.name = "Calibri"
    table_header.font.size = Pt(9.5)
    table_header.font.bold = True
    table_header.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    table_header.paragraph_format.space_before = Pt(0)
    table_header.paragraph_format.space_after = Pt(0)
    table_header.paragraph_format.line_spacing = 1.05

    callout = styles["Callout"]
    callout.font.name = "Calibri"
    callout.font.size = Pt(10.5)
    callout.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    callout.paragraph_format.space_before = Pt(0)
    callout.paragraph_format.space_after = Pt(0)
    callout.paragraph_format.line_spacing = 1.2

    small = styles["Small Note"]
    small.font.name = "Calibri"
    small.font.size = Pt(9)
    small.font.color.rgb = RGBColor.from_string(MUTED)
    small.paragraph_format.space_before = Pt(0)
    small.paragraph_format.space_after = Pt(4)
    small.paragraph_format.line_spacing = 1.1


def set_core_properties(doc):
    props = doc.core_properties
    props.title = "STRH HRIS Technical and Implementation Proposal"
    props.subject = "TOR-aligned proposal for the development, customization, deployment, and support of an HRIS"
    props.author = "[INSERT BIDDER / SERVICE PROVIDER NAME]"
    props.keywords = "STRH, HRIS, technical proposal, implementation, DOH"
    props.comments = "Draft for review. Replace all bracketed fields before submission."


def setup_sections(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [5400, 3960], indent=0)
    set_table_borders(table, color=WHITE, size="0")
    left, right = table.rows[0].cells
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("STRH HRIS TECHNICAL PROPOSAL")
    set_run_font(r, size=8.5, bold=True, color=MUTED)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("DRAFT FOR REVIEW")
    set_run_font(r, size=8.5, bold=True, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Page ")
    set_run_font(r, size=8.5, color=MUTED)
    field_run = add_field(p, "PAGE")
    set_run_font(field_run, size=8.5, color=MUTED)

    first_footer = section.first_page_footer
    p = first_footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Draft for Review  |  18 August 2026")
    set_run_font(r, size=8.5, color=MUTED)


def add_para(doc, text="", *, style=None, size=None, bold=False, italic=False, color=None,
             align=None, before=0, after=8, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_with_next = True
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    apply_num(p, 90)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph()
    apply_num(p, 91)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_callout(doc, label, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color="C9D8EB", size="6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.style = doc.styles["Callout"]
    label_run = p.add_run(f"{label}: ")
    set_run_font(label_run, size=10.5, bold=True, color=DARK_BLUE)
    text_run = p.add_run(text)
    set_run_font(text_run, size=10.5, color=DARK_BLUE)
    add_para(doc, after=3)


def fill_table(table, headers, rows, widths, aligns=None, status_column=None):
    set_table_geometry(table, widths)
    set_table_borders(table)
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.style = table._parent.part.document.styles["Table Header"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if aligns and aligns[idx] == "center" else WD_ALIGN_PARAGRAPH.LEFT
        p.add_run(str(text))
    set_repeat_table_header(table.rows[0])
    for r_idx, row_values in enumerate(rows, start=1):
        row = table.rows[r_idx]
        for c_idx, value in enumerate(row_values):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            p.style = table._parent.part.document.styles["Table Text"]
            if aligns and aligns[c_idx] == "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(str(value))
            if status_column is not None and c_idx == status_column:
                status = str(value).lower()
                if "available" in status:
                    set_cell_shading(cell, "EAF5EE")
                elif "included" in status:
                    set_cell_shading(cell, LIGHT_BLUE)
                elif "dependency" in status:
                    set_cell_shading(cell, PALE_GOLD)
                elif "clarification" in status:
                    set_cell_shading(cell, PALE_RED)
    return table


def add_table(doc, headers, rows, widths, aligns=None, status_column=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    fill_table(table, headers, rows, widths, aligns=aligns, status_column=status_column)
    add_para(doc, after=2)
    return table


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_cover(doc):
    add_para(doc, before=8, after=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run()
    run.add_picture(str(LOGO), width=Inches(1.35))
    doc_pr = run._r.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[0].set("descr", "Southern Tagalog Regional Hospital seal")

    add_para(doc, "TECHNICAL AND IMPLEMENTATION PROPOSAL", style="Proposal Kicker",
             align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    p = add_para(doc, "Human Resource Information System", style="Title",
                 align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    for run in p.runs:
        set_run_font(run, size=28, bold=True, color=BLACK)
    add_para(doc, "Development, Customization, Deployment, Training, and Support",
             style="Subtitle", align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_para(doc, "Prepared for DOH - Southern Tagalog Regional Hospital",
             size=11, bold=True, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=20)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(14)
    set_paragraph_border(rule, color=GOLD, size=18)

    table = doc.add_table(rows=4, cols=2)
    set_table_geometry(table, [4680, 4680])
    set_table_borders(table, color="D7DEE8", size="6")
    data = [
        ("Prepared by", "[INSERT BIDDER / SERVICE PROVIDER NAME]"),
        ("Authorized representative", "[INSERT NAME AND POSITION]"),
        ("Proposal date", "18 August 2026"),
        ("Proposal reference", "[INSERT REFERENCE NUMBER]"),
    ]
    for i, (label, value) in enumerate(data):
        left, right = table.rows[i].cells
        set_cell_shading(left, LIGHT)
        p = left.paragraphs[0]
        p.style = doc.styles["Table Header"]
        p.add_run(label)
        p = right.paragraphs[0]
        p.style = doc.styles["Table Text"]
        p.add_run(value)

    add_para(doc, after=18)
    add_para(doc,
             "DRAFT FOR REVIEW - Replace all bracketed fields and reconcile this proposal with the final bidding documents, bid bulletins, approved budget, and commercial terms before submission.",
             size=9, bold=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)


def build_document():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    setup_numbering(doc)
    setup_sections(doc)
    set_core_properties(doc)
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    add_cover(doc)
    add_page_break(doc)

    add_heading(doc, "Proposal Overview", 1)
    add_callout(doc, "Proposal position",
                "The existing HRIS is a substantial implementation baseline. This proposal distinguishes capabilities that are already available from functions that will be completed during the project, items that depend on STRH inputs, and matters requiring formal clarification. No unfinished feature is represented as currently operational.")
    add_table(doc,
              ["Document item", "Draft value"],
              [
                  ("Procuring entity", "DOH - Southern Tagalog Regional Hospital"),
                  ("Project", "Development of a Human Resource Information System"),
                  ("Proposal type", "Technical and Implementation Proposal"),
                  ("Implementation period", "Thirty-two (32) calendar weeks from Notice to Proceed, subject to approved dependencies"),
                  ("Support / warranty", "To be finalized against the bidding documents and commercial schedule"),
                  ("Commercial validity", "[INSERT VALIDITY PERIOD]"),
              ],
              [2700, 6660])

    add_heading(doc, "Status Definitions", 2)
    add_table(doc,
              ["Status", "Meaning in this proposal"],
              [
                  ("Available baseline", "Implemented in the current solution and available for demonstration, subject to configuration, data validation, security review, and UAT."),
                  ("Included for completion", "Not yet complete. Included as a project deliverable with a milestone and acceptance criterion."),
                  ("STRH dependency", "Completion requires official data, forms, policies, signatories, infrastructure, devices, access, or a documented STRH decision."),
                  ("Clarification required", "The TOR or final bidding documents must resolve the requirement before a compliance commitment, price, or acceptance method is finalized."),
                  ("Optional enhancement", "Excluded from the base offer unless expressly selected and priced."),
              ],
              [2500, 6860], status_column=0)

    add_page_break(doc)
    add_heading(doc, "Contents", 1)
    for item in [
        "1. Executive Summary",
        "2. Understanding of the Requirement",
        "3. Proposed HRIS Scope",
        "4. Technical Architecture and Security",
        "5. Implementation Methodology and Schedule",
        "6. Data Migration, Integration, and Deployment",
        "7. Testing, Acceptance, and Quality Assurance",
        "8. Training, Documentation, Support, and Warranty",
        "9. Deliverables and Acceptance Matrix",
        "10. Assumptions, Dependencies, Clarifications, and Exclusions",
        "11. Commercial Proposal",
        "12. Proposal Acceptance and Authorization",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Document Control", 2)
    add_table(doc,
              ["Control", "Draft value"],
              [
                  ("Version", "0.1"),
                  ("Status", "Draft for review and completion"),
                  ("Date", "18 August 2026"),
                  ("Proposal owner", "[INSERT BIDDER / SERVICE PROVIDER NAME]"),
                  ("Primary basis", "STRH Terms of Reference and current HRIS implementation baseline"),
              ],
              [2500, 6860])

    add_page_break(doc)
    add_heading(doc, "1. Executive Summary", 1)
    add_para(doc,
             "[INSERT BIDDER / SERVICE PROVIDER NAME] respectfully submits this proposal for the development, customization, installation, configuration, deployment, training, and support of a secure Human Resource Information System for DOH - Southern Tagalog Regional Hospital (STRH). The proposed engagement uses the existing HRIS implementation as an accelerator while retaining a controlled requirements-validation, completion, testing, and acceptance process.")
    add_para(doc,
             "The proposed solution is designed to centralize employee information, digitize 201 files, manage Plantilla and personnel movements, process attendance and leave, integrate biometric devices, provide employee self-service, generate management information, and strengthen auditability and access control. The project also covers deployment preparation, data migration, documentation, training, warranty, and support activities required for operational adoption.")
    add_para(doc,
             "The current baseline already contains major HR workflows. Several TOR requirements remain to be completed or formally validated. These items are presented as scheduled deliverables rather than as existing capabilities. Final compliance will be demonstrated through agreed acceptance criteria, STRH-approved templates and policies, role-based UAT, deployment validation, and documented sign-off.")

    add_heading(doc, "Proposed outcomes", 2)
    for text in [
        "A centralized, searchable, and controlled employee information repository with major 201-file records.",
        "Integrated Plantilla, personnel movement, service history, attendance, leave, and employee self-service workflows.",
        "Biometric-assisted DTR processing with exception handling, correction controls, locking, and audit records.",
        "Configurable roles and permissions, secure sessions, action logging, and controlled document generation.",
        "STRH-prioritized reports, official forms, migration controls, backup and recovery procedures, training, and technical documentation.",
        "A phased implementation with transparent status reporting and measurable acceptance gates.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "Recommended contracting principle", 2)
    add_callout(doc, "Truthful feature commitment",
                "A feature may be included even when it is not yet finished, provided it is expressly identified as an implementation deliverable, included in the price and schedule, supported by acceptance criteria, and not described as currently available.", fill=PALE_GOLD)

    add_heading(doc, "2. Understanding of the Requirement", 1)
    add_para(doc,
             "STRH requires more than an employee database. The TOR describes an integrated institutional HR platform together with implementation services, biometric hardware and integration, data migration, quality assurance, training, documentation, support, security, backup, archiving, and recovery obligations. The proposal therefore treats operational readiness and adoption as contract deliverables alongside software features.")

    add_heading(doc, "Business objectives", 2)
    for text in [
        "Replace fragmented and paper-heavy HR processes with a controlled digital system.",
        "Improve the completeness, accuracy, timeliness, and traceability of employee information.",
        "Support attendance, leave, personnel actions, reporting, employee services, and management decision-making.",
        "Align the system with applicable CSC, DBM, DOH, privacy, records-management, and institutional policies.",
        "Provide a maintainable solution that STRH administrators and IT personnel can operate and support.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "Scope interpretation requiring confirmation", 2)
    add_table(doc,
              ["TOR matter", "Proposal treatment"],
              [
                  ("Payroll processing", "The TOR objective mentions payroll, while the detailed module list focuses on HR, attendance, leave, and payroll notifications. Full payroll computation and disbursement require clarification and separate acceptance rules."),
                  ("Recruitment", "Recruitment appears in objectives, forms, and statistics, but a complete applicant-tracking workflow is not fully specified. Confirm whether it is a base module or future phase."),
                  ("Technology framework", "The current solution uses React/TanStack/Vite, a Node.js API, and MySQL. Written acceptance of this JavaScript-based stack should be obtained where the TOR language is interpreted as requiring Laravel or .NET."),
                  ("Database compatibility", "The current implementation uses MySQL. The proposal does not claim Oracle, PostgreSQL, MSSQL, or DBF portability unless separately engineered and tested."),
                  ("Legacy browsers", "A modern browser support matrix is proposed. Internet Explorer and unspecified legacy-browser compatibility require a formal exception or a separately funded compatibility approach."),
                  ("Support term", "The TOR refers to both one-year maintenance and two-year proactive assistance. The final duration, inclusions, and pricing must be reconciled in the contract."),
              ],
              [2550, 6810])

    add_heading(doc, "3. Proposed HRIS Scope", 1)
    add_para(doc,
             "The following scope is grouped by capability and proposal status. Available baseline functions remain subject to STRH configuration, official data validation, security review, and UAT. Items marked Included for completion are project commitments, not representations of current readiness.")

    add_heading(doc, "3.1 Employee Information and 201 Files", 2)
    add_table(doc,
              ["Capability", "Status", "Proposed delivery and acceptance"],
              [
                  ("Employee master list and profile", "Available baseline", "Search, filter, create, update, archive/restore, and controlled profile access; validate with STRH master data and permission scenarios."),
                  ("Major 201-file sections", "Available baseline", "Personal, family, education, eligibility, work, training, organization, salary, service, IPCR, and related records; complete official field mapping during validation."),
                  ("Reference libraries and organizational hierarchy", "Available baseline", "Departments, positions, salary schedules, HR references, and hierarchy management; populate official STRH codes and effective dates."),
                  ("Official PDS and WES generation", "Available baseline", "Existing export paths will be validated for official field coverage, overflow, continuation sheets, permissions, and representative render samples."),
                  ("Administrative and disciplinary cases", "Included for completion", "Confidential case register, nature of case, sanction, dates, remarks, attachments, restricted permissions, audit events, and authorized reporting."),
                  ("Qualification standards and evaluation", "Included for completion", "Define the STRH-approved qualification standard model and evaluation rules before implementation and UAT."),
              ],
              [2900, 1750, 4710], aligns=["left", "center", "left"], status_column=1)

    add_heading(doc, "3.2 Plantilla, Personnel Movements, and Service Records", 2)
    add_table(doc,
              ["Capability", "Status", "Proposed delivery and acceptance"],
              [
                  ("Plantilla / PSIPOP management", "Available baseline", "Authorized items, funding and salary references, occupancy, vacancy, lifecycle, history, and reconciliation; validate against STRH PSIPOP source."),
                  ("Personnel movement workflow", "Available baseline", "Draft, submit, review, approve, post, reject, reverse, authority references, conflict controls, and event history; confirm action-specific rules."),
                  ("Appointment and personnel action coverage", "Available baseline", "Supports major TOR action categories; action-specific mapping, documents, and separation-of-duties require STRH UAT."),
                  ("Service record generation", "Available baseline", "Generated from posted movements and controlled legacy entries; official STRH template and signatory rules remain dependencies."),
                  ("Secure supporting documents", "Included for completion", "Controlled attachment storage, file validation, access authorization, retention, and audit trail for movements and personnel records."),
                  ("Payroll effect notification", "Included for completion", "Notify designated payroll recipients for approved salary, step, loyalty, or related actions, with acknowledgement tracking."),
              ],
              [2900, 1750, 4710], aligns=["left", "center", "left"], status_column=1)

    add_heading(doc, "3.3 Attendance, DTR, and Biometric Integration", 2)
    add_table(doc,
              ["Capability", "Status", "Proposed delivery and acceptance"],
              [
                  ("Attendance imports and DTR processing", "Available baseline", "Import, refresh, DTR calculation, manual entries, Excel/PDF output, and bounded date processing; validate schedules, holidays, rest days, overnight duty, and time zone."),
                  ("Manual corrections and field locking", "Available baseline", "Employee requests, approval/reversal, per-time-slot locks, immutable events, and audit records; validate concurrent edits and biometric refresh."),
                  ("Schedules and date-specific overrides", "Available baseline", "Shift templates, assignments, bulk scheduling, and exceptions; validate hospital schedule patterns and alternative work arrangements."),
                  ("Biometric devices and ADMS", "Available baseline", "Device configuration, connectivity, ADMS endpoints, synchronization, real-time status/logs, mapping exceptions, and retry; actual device UAT is an STRH dependency."),
                  ("Online alternative attendance", "Included for completion", "Implement only after STRH approves the eligible work arrangements, authentication, geolocation/device controls, and approval policy."),
                  ("Travel, OB, overtime, offsetting, and CTO codes", "Included for completion", "Create configurable activity-code and workflow rules with schedule, pay, DTR, approval, and reporting effects."),
              ],
              [2900, 1750, 4710], aligns=["left", "center", "left"], status_column=1)

    add_heading(doc, "3.4 Leave Management", 2)
    add_table(doc,
              ["Capability", "Status", "Proposed delivery and acceptance"],
              [
                  ("Leave types, filing, balances, and ledger", "Available baseline", "Configurable leave metadata, applications, balance checks, approval deduction, adjustments, ledger, and reversal; validate with official HR leave cards."),
                  ("Form 6 generation", "Available baseline", "Excel/PDF generation with STRH-approved signatories, field mapping, dependencies, and representative samples."),
                  ("Multi-level recommendation, certification, and final action", "Included for completion", "Configurable routing, independent credit certification, immediate-superior action, delegation, signatories, separation of duties, and complete event history."),
                  ("Automatic leave-credit accrual", "Included for completion", "Effective-dated policy engine with preview, validation, idempotent posting, ledger reconciliation, approval, and rerun protection."),
                  ("Low-balance and payroll-risk warnings", "Included for completion", "Configurable thresholds, recipients, deduplication, acknowledgement, and reports."),
                  ("Offline leave capture", "STRH dependency", "Confirm whether authorized HR encoding of signed applications satisfies the requirement or whether document capture and source tracking are required."),
              ],
              [2900, 1750, 4710], aligns=["left", "center", "left"], status_column=1)

    add_heading(doc, "3.5 Employee Self-Service and Portal", 2)
    add_table(doc,
              ["Capability", "Status", "Proposed delivery and acceptance"],
              [
                  ("Employee login and own-record access", "Available baseline", "Linked employee accounts, profile access, leave filing/history, attendance, DTR corrections, requests, and PDS generation; validate privacy boundaries."),
                  ("Personal-data maintenance", "Available baseline", "Approved personal/contact and designated self-service sections only; HR-controlled employment and salary data remain governed."),
                  ("Overtime, CTO, OB, certificate, and schedule-change requests", "Included for completion", "Replace current placeholders with routed, auditable workflows, attachments where approved, employee actions, notifications, and unified request history."),
                  ("Announcements, policies, and news", "Included for completion", "Authorized publishing, effective dates, visibility groups, attachments, expiration, and audit controls."),
                  ("Official email notifications", "STRH dependency", "Requires approved mail service, sender domain, credentials, network access, templates, and recipient rules."),
              ],
              [2900, 1750, 4710], aligns=["left", "center", "left"], status_column=1)

    add_heading(doc, "3.6 Reports, Analytics, and Forms", 2)
    add_para(doc,
             "The current baseline includes dashboard analytics and a personnel/Plantilla report family. The project will prioritize reports through an approved catalog rather than claiming that every listed report already exists.")
    for text in [
        "Employee master list, demographic profiles, personnel statistics, filled and unfilled positions, and organization distributions.",
        "Attendance, tardiness, undertime, absences, overtime, schedules, travel, OB, perfect attendance, and DTR summaries.",
        "Leave balances, ledger cards, certifications, forced leave, late-filed leave, disapproved applications, sick-leave profiles, and summaries.",
        "Plantilla, service record, salary adjustment, step increment, loyalty, COE, and other employee certifications.",
        "Saved filters and controlled report definitions after standard reports and official output formats are accepted.",
    ]:
        add_bullet(doc, text)
    add_callout(doc, "Report acceptance",
                "Each report will have an approved definition covering purpose, authorized users, data scope, filters, calculations, as-of date, totals, output format, record count, and sample sign-off.")

    add_heading(doc, "3.7 Administration, Security, Backup, and Recovery", 2)
    add_table(doc,
              ["Capability", "Status", "Proposed delivery and acceptance"],
              [
                  ("User accounts, roles, and action permissions", "Available baseline", "Unlimited account creation subject to capacity, configurable role-permission matrix, employee linking, lock/unlock, reset, and deny-by-default enforcement."),
                  ("Audit and error logs", "Available baseline", "Action and error history; complete retained-history pagination, filters, controlled export, retention, and monitoring alerts during implementation."),
                  ("Configurable session timeout", "Included for completion", "Administrator-configurable duration within approved security limits, with documented idle/absolute timeout behavior."),
                  ("Complete automated backups", "Included for completion", "Encrypted daily and monthly backup, signed manifest, off-system copy, retention, failure alerts, and restricted download/access."),
                  ("Restore and recovery testing", "Included for completion", "Controlled restore tool/runbook, 12-hour recovery objective validation, periodic witnessed restore test, and evidence package."),
                  ("Archiving and reload", "Included for completion", "Policy-aligned archival package, integrity checks, authorized reload, retention, and destruction workflow subject to STRH/NAP policy."),
              ],
              [2900, 1750, 4710], aligns=["left", "center", "left"], status_column=1)

    add_heading(doc, "4. Technical Architecture and Security", 1)
    add_heading(doc, "4.1 Proposed application architecture", 2)
    add_para(doc,
             "The proposed solution uses a separated web user interface, application/API service, and relational database. The current implementation uses React 19 with TanStack Router and Vite for the browser application, a Node.js API service for server-side business rules and authorization, Python tooling for prescribed document generation, and MySQL 8 for persistent data.")
    add_table(doc,
              ["Layer", "Proposed technology and responsibility"],
              [
                  ("Presentation", "Responsive browser interface for supported desktop and mobile browsers; role- and permission-aware navigation and actions."),
                  ("Application/API", "Node.js service enforcing authentication, permissions, workflow rules, validation, audit events, integrations, and document-job authorization."),
                  ("Database", "MySQL 8 relational schema with constraints, indexes, migrations, backup, restore, and recovery procedures."),
                  ("Document generation", "Controlled Python/LibreOffice-compatible generation for approved forms and reports, with render and sample validation."),
                  ("Integration", "Biometric ADMS/TCP/IP, uploaded attendance files, email service where approved, and future interfaces governed by documented contracts."),
                  ("Deployment", "Windows Server 2019+ or an approved 64-bit open-source server OS, behind HTTPS and a controlled reverse proxy."),
              ],
              [2300, 7060])

    add_heading(doc, "4.2 Security and privacy controls", 2)
    for text in [
        "HTTPS/TLS for all production access, secure cookies, approved origins, and restricted administrative endpoints.",
        "Server-side authorization for every protected action, with role-permission mapping and record ownership checks.",
        "Three-attempt account lock, password change/reset controls, password history, session expiry, and administrator auditability.",
        "Data minimization, least privilege, protected HR documents, controlled exports, retention, and accountable access to sensitive records.",
        "Audit trails for authentication, record changes, approvals, corrections, exports, and administrative actions.",
        "Encrypted and isolated backup storage, secrets management, logging, monitoring, dependency review, and vulnerability remediation before go-live.",
    ]:
        add_bullet(doc, text)
    add_callout(doc, "Production security gate",
                "Production acceptance requires deployment hardening, valid SSL, permission testing, backup/restore evidence, vulnerability review, and closure of all agreed critical and high-risk findings.", fill=PALE_RED)

    add_heading(doc, "4.3 Supported client and server environment", 2)
    add_table(doc,
              ["Environment", "Proposed support position"],
              [
                  ("Desktop browsers", "Current stable Microsoft Edge, Google Chrome, and Mozilla Firefox on supported Windows versions."),
                  ("Mobile browsers", "Current stable Chrome on Android and Safari on supported iOS versions; formal device matrix to be approved."),
                  ("Legacy browser", "Internet Explorer and unspecified obsolete browsers are not assumed. Clarification or exception is required."),
                  ("Server OS", "Windows Server 2019 or later, or an approved 64-bit open-source server OS after deployment certification."),
                  ("Database", "MySQL 8 is the proposed and currently implemented database. Other database engines are excluded unless separately scoped."),
              ],
              [2300, 7060])

    add_heading(doc, "5. Implementation Methodology and Schedule", 1)
    add_para(doc,
             "A thirty-two-week implementation is proposed because the engagement includes requirements validation, completion of material TOR gaps, official-data and form reconciliation, infrastructure and biometric integration, testing, training, and controlled deployment. Activities may overlap only where dependencies and quality gates allow.")

    add_table(doc,
              ["Phase", "Indicative timing", "Principal outputs / gate"],
              [
                  ("0. Mobilization and kickoff", "Week 1", "Governance, communication plan, access, issue log, source baseline, security rules, and detailed work plan."),
                  ("1. Requirements validation and Inception Report", "Weeks 2-4", "Approved scope, compliance matrix, workflows, data dictionary, report catalog, architecture, dependencies, acceptance criteria, and change control."),
                  ("2. STRH configuration and data preparation", "Weeks 5-9", "Official reference values, roles, signatories, initial migration mapping, data-quality report, templates, and environment readiness."),
                  ("3. Completion of priority TOR workflows", "Weeks 10-19", "Disciplinary cases, self-service requests, approval routing, activity codes, leave automation, attachments, notifications, and controlled demonstrations."),
                  ("4. Reports, security, backup, and operations", "Weeks 20-26", "Prioritized reports/forms, session configuration, log operations, automated backup, restore tooling, archival controls, and runbooks."),
                  ("5. Integration, migration rehearsal, and UAT", "Weeks 27-30", "Device/network validation, migration rehearsal, role-based UAT, defect closure, render samples, training, and readiness recommendation."),
                  ("6. Production deployment and stabilization", "Weeks 31-32", "Approved production installation, final migration, smoke tests, go-live, hypercare, acceptance evidence, and turnover."),
              ],
              [1900, 1600, 5860], aligns=["left", "center", "left"])

    add_heading(doc, "Project governance", 2)
    for text in [
        "Weekly implementation meeting and concise status report covering progress, next actions, risks, decisions, dependencies, and change requests.",
        "A single approved requirements baseline and change-control register after the Inception Report.",
        "Named STRH product owner, HR process owners, IT/security representatives, data owners, UAT leads, and signatories.",
        "Formal gates for scope approval, migration readiness, UAT entry, production readiness, go-live, and final acceptance.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "6. Data Migration, Integration, and Deployment", 1)
    add_heading(doc, "6.1 Data migration", 2)
    add_numbered(doc, "Inventory source files, database extracts, records, codes, attachments, owners, sensitivity, and retention requirements.")
    add_numbered(doc, "Approve the mapping specification, normalization rules, duplicate strategy, mandatory fields, and rejection handling.")
    add_numbered(doc, "Run a non-production trial migration and issue reconciliation totals, rejected-row reports, and representative record checks.")
    add_numbered(doc, "Obtain HR validation, correct source/mapping issues, and repeat until acceptance thresholds are met.")
    add_numbered(doc, "Execute final migration during the approved cutover window, validate control totals, and obtain signed migration acceptance.")

    add_heading(doc, "6.2 Biometric integration and hardware", 2)
    add_para(doc,
             "The software baseline contains biometric device configuration, ADMS integration paths, synchronization, device monitoring, and attendance exception handling. Final acceptance requires testing with the actual contracted devices, STRH network segments, employee device identifiers, time settings, outages, retries, and reconciliation procedures.")
    add_table(doc,
              ["Item", "Proposal treatment"],
              [
                  ("Five biometric units", "Included only when specified and priced in the final Bill of Quantities; model/equivalency, warranties, accessories, and delivery must be confirmed."),
                  ("Cabling and network", "STRH and bidder responsibilities, ports, VLAN/firewall rules, Wi-Fi/TCP/IP, power, installation points, and testing must be documented."),
                  ("Employee mapping", "STRH will provide or validate the mapping between HR employee identifiers and biometric user identifiers."),
                  ("Exception reconciliation", "Unmatched or invalid punches will be quarantined, reviewed, mapped, reprocessed, and reported rather than silently discarded."),
              ],
              [2500, 6860])

    add_heading(doc, "6.3 Production deployment", 2)
    for text in [
        "Approved production topology, server sizing, operating system, database, reverse proxy, domain, SSL certificate, firewall, storage, and administrative access.",
        "Versioned configuration and secrets management without embedding production credentials in source code or documents.",
        "Deterministic database migrations, deployment checklist, health checks, export dependency checks, rollback plan, and configuration backup.",
        "Monitoring for availability, errors, backup outcomes, storage, resource usage, and security events.",
        "Cutover, smoke testing, go-live authorization, and stabilization support with named escalation contacts.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "7. Testing, Acceptance, and Quality Assurance", 1)
    add_para(doc,
             "Build success alone is not acceptance. The project will use traceable requirements, controlled test data, role-based scenarios, output validation, data reconciliation, security checks, and documented defect closure.")

    add_table(doc,
              ["Test level", "Minimum coverage"],
              [
                  ("Automated regression", "Authentication, permissions, employee ownership, attendance locks/import exceptions, leave ledger math, personnel movement posting/reversal, migration checks, exports, and backup/restore controls."),
                  ("System integration", "Database, biometric devices, network, email where approved, file imports, document generation, and production dependencies."),
                  ("Security and privacy", "Role/action matrix, direct API authorization, session behavior, sensitive exports, attachments, audit coverage, configuration, vulnerability review, and backup access."),
                  ("Migration validation", "Source-to-target counts, required fields, duplicates, rejected records, reference mapping, balances, history, and signed samples."),
                  ("User acceptance", "HR, Approver, Administrator, Viewer, Employee, and IT scenarios using approved data, forms, devices, calculations, reports, and exception cases."),
                  ("Operational readiness", "Installation, restart, monitoring, backup, restore, fail/retry, troubleshooting, escalation, rollback, and knowledge-transfer exercises."),
              ],
              [2500, 6860])

    add_heading(doc, "Acceptance principle", 2)
    add_callout(doc, "Definition of done",
                "A requirement is complete only when its agreed implementation, permission controls, validation, audit behavior, output, documentation, and UAT evidence are accepted. A screen, route, or placeholder is not sufficient evidence of completion.")

    add_heading(doc, "8. Training, Documentation, Support, and Warranty", 1)
    add_heading(doc, "8.1 Training and knowledge transfer", 2)
    add_table(doc,
              ["Audience", "Proposed coverage"],
              [
                  ("STRH IT Division", "Architecture, deployment, configuration, database basics, backup/restore, monitoring, security, troubleshooting, logs, escalation, and controlled maintenance."),
                  ("Global / Super Administrator", "Users, roles, permissions, settings, references, audit, errors, accounts, sessions, backup operations, and security responsibilities."),
                  ("HRD System Administrator and HR users", "Employee/201, Plantilla, movements, service records, attendance, leave, approvals, reports, correction/reversal, and operational controls."),
                  ("Approvers and supervisors", "Assigned approvals, recommendations, certification, decisions, delegation, notifications, and accountability."),
                  ("End users", "Four batches covering login, profile, self-service, attendance, leave, requests, forms, password management, privacy, and issue reporting."),
              ],
              [2800, 6560])

    add_heading(doc, "8.2 Documentation set", 2)
    for text in [
        "Approved Inception Report and requirements/compliance matrix.",
        "User manual, operations manual, administrator manual, technical documentation, policies/procedures, FAQ, and quick-reference guides.",
        "Architecture, data dictionary, integration specifications, configuration register, migration mapping, report catalog, and deployment runbook.",
        "Backup/restore, monitoring, troubleshooting, escalation, release, rollback, retention, and security procedures.",
        "Test plan, UAT scripts, defect register, training records, acceptance evidence, and turnover inventory.",
        "Three hard and soft copies where required by the final contract, with controlled electronic source files.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "8.3 Support and service levels", 2)
    add_table(doc,
              ["Severity", "Target response", "Target resolution / treatment"],
              [
                  ("Critical - system unavailable", "Within 4 hours", "Within 24 hours where technically feasible; continuous escalation, workaround and status updates."),
                  ("High - major function impaired", "Within 1 business day", "Within 3 business days, subject to dependency and approved workaround."),
                  ("Low - minor issue or request", "Within 3 business days", "Within 7 business days or scheduled release, as appropriate."),
              ],
              [2850, 1950, 4560], aligns=["left", "center", "left"])
    add_para(doc,
             "Final SLA clocks, exclusions, support channels, office hours, on-site response, service credits or remedies, maintenance windows, warranty duration, and post-warranty pricing must be reconciled with the final bidding documents and contract.", style="Small Note")

    add_heading(doc, "9. Deliverables and Acceptance Matrix", 1)
    add_table(doc,
              ["Deliverable", "Target phase", "Acceptance evidence"],
              [
                  ("Approved Inception Report", "Phase 1", "Signed scope, workflows, architecture, compliance matrix, acceptance criteria, schedule, dependencies, and change control."),
                  ("Configured HRIS baseline", "Phase 2", "Approved references, roles, signatories, branding, forms, initial data, and demonstration results."),
                  ("Completed priority TOR workflows", "Phase 3", "Feature demonstrations, permission tests, event/audit evidence, approved outputs, and closed material defects."),
                  ("Reports and prescribed outputs", "Phase 4", "Approved definitions, representative files, totals, filters, signatories, render checks, and user sign-off."),
                  ("Security and operational controls", "Phase 4", "Hardening checklist, permission matrix, session rules, monitoring, encrypted backups, restore test, and recovery evidence."),
                  ("Migration package", "Phases 2 and 5", "Mapping, scripts/tools, rejected-row report, reconciliation totals, trial results, final validation, and HR sign-off."),
                  ("Biometric and integration validation", "Phase 5", "Device inventory, connection results, mapped users, synchronized logs, exception reconciliation, outage/retry tests, and sign-off."),
                  ("UAT and quality package", "Phase 5", "Test scripts, results, defects, retests, security checks, output samples, and acceptance recommendation."),
                  ("Training and documentation", "Phases 5 and 6", "Approved manuals, materials, attendance, evaluations, recordings where allowed, and knowledge-transfer sign-off."),
                  ("Production deployment and turnover", "Phase 6", "Installation record, configuration inventory, smoke test, cutover/recovery plan, source/deployment package, and go-live acceptance."),
              ],
              [3300, 1500, 4560], aligns=["left", "center", "left"])

    add_heading(doc, "10. Assumptions, Dependencies, Clarifications, and Exclusions", 1)
    add_heading(doc, "10.1 STRH responsibilities and dependencies", 2)
    for text in [
        "Provide an authorized project owner, HR process owners, IT/security representatives, data owners, UAT users, and decision-makers.",
        "Provide official employee data, PSIPOP/Plantilla source, reference codes, policies, schedules, leave rules, reports, forms, signatories, branding, and retention requirements in agreed formats and dates.",
        "Provide access to approved servers, network, domains, SSL process, email service, biometric devices, installation locations, and required technical personnel.",
        "Review deliverables and issue consolidated decisions within the agreed review period; delays may move dependent milestones.",
        "Authorize and protect test data, accounts, devices, and production credentials in accordance with STRH security and privacy requirements.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "10.2 Required pre-submission or inception clarifications", 2)
    add_table(doc,
              ["Clarification", "Required decision"],
              [
                  ("Payroll", "Full payroll module, payroll interface/notification only, or excluded from the base scope."),
                  ("Recruitment", "Full recruitment lifecycle or limited forms, records, reports, and personnel movement support."),
                  ("Technology", "Formal acceptance of the proposed JavaScript/Node/MySQL architecture."),
                  ("Browser support", "Approved modern browser/device matrix and treatment of Internet Explorer/legacy systems."),
                  ("Database scope", "MySQL as the production database or funded portability to additional engines."),
                  ("Biometric supply", "Exact technical specification, equivalency, unit count, accessories, warranties, cabling, installation, and acceptance."),
                  ("Official templates", "Editable source forms, report formats, signatory rules, continuation sheets, and approval authority."),
                  ("Warranty/support", "One-year maintenance and two-year proactive assistance language, coverage, on-site obligations, and pricing."),
                  ("Ownership and licensing", "Source-code handover, license, reuse rights, third-party components, escrow if required, and post-contract maintenance rights."),
              ],
              [2800, 6560])

    add_heading(doc, "10.3 Base exclusions unless expressly included", 2)
    for text in [
        "Full payroll computation, remittance, disbursement, and accounting integration until the scope is clarified and priced.",
        "Native Android/iOS applications; the base solution is responsive web and mobile-browser compatible.",
        "Compatibility with Oracle, PostgreSQL, MSSQL, DBF, Internet Explorer, or unsupported legacy operating systems.",
        "Third-party subscriptions, SMS charges, email-service fees, SSL/domain fees, cloud hosting, operating-system licenses, or database licenses not expressly listed in the Bill of Quantities.",
        "Data cleansing beyond agreed transformation rules, manual reconstruction of missing records, or scanning/indexing of legacy paper files unless separately quantified.",
        "New requirements introduced after approval of the Inception Report, except through approved change control.",
    ]:
        add_bullet(doc, text)

    add_page_break(doc)
    add_heading(doc, "11. Commercial Proposal", 1)
    add_callout(doc, "Commercial completion required",
                "The amounts below are deliberately left blank. They must be completed by the bidder after final scope, hardware, taxes, warranty, licensing, and procurement requirements are confirmed.", fill=PALE_GOLD)
    add_heading(doc, "Commercial cost schedule", 2)
    add_table(doc,
              ["Cost component", "Amount (PHP)"],
              [
                  ("Software implementation, customization, and configuration", "[INSERT AMOUNT]"),
                  ("Data migration and validation", "[INSERT AMOUNT]"),
                  ("Reports, forms, and document templates", "[INSERT AMOUNT]"),
                  ("Biometric hardware - five units, accessories, delivery, and installation", "[INSERT AMOUNT]"),
                  ("Deployment infrastructure and third-party licenses, if included", "[INSERT AMOUNT]"),
                  ("Training, documentation, and knowledge transfer", "[INSERT AMOUNT]"),
                  ("Warranty, maintenance, and support", "[INSERT AMOUNT]"),
                  ("Applicable taxes", "[INSERT AMOUNT]"),
                  ("TOTAL PROPOSAL PRICE", "[INSERT TOTAL]"),
              ],
              [6800, 2560], aligns=["left", "center"])

    add_heading(doc, "Proposed payment milestones", 2)
    add_table(doc,
              ["Milestone", "Indicative share", "Payment evidence"],
              [
                  ("Contract mobilization and approved Inception Report", "[INSERT %]", "Approved deliverable and invoice requirements."),
                  ("Configured baseline and accepted migration rehearsal", "[INSERT %]", "Phase-gate approval and reconciliation report."),
                  ("Completion of priority TOR workflows", "[INSERT %]", "Demonstration and accepted test evidence."),
                  ("UAT completion and production readiness", "[INSERT %]", "Signed UAT and readiness recommendation."),
                  ("Go-live, turnover, and final acceptance", "[INSERT %]", "Production acceptance and complete turnover package."),
              ],
              [4450, 1600, 3310], aligns=["left", "center", "left"])
    add_para(doc,
             "Payment terms are subject to the final bidding documents, government accounting and auditing rules, applicable taxes, performance security, warranty security, and the executed contract.", style="Small Note")

    add_heading(doc, "12. Proposal Acceptance and Authorization", 1)
    add_para(doc,
             "This draft is submitted for review and completion. It does not become a binding offer until all bracketed fields, technical qualifications, schedules, prices, securities, signatures, and required procurement forms have been completed and authorized by the bidder.")

    add_heading(doc, "Submitted by", 2)
    add_table(doc,
              ["Field", "Authorized bidder information"],
              [
                  ("Company / service provider", "[INSERT LEGAL NAME]"),
                  ("Authorized representative", "[INSERT NAME]"),
                  ("Position", "[INSERT POSITION]"),
                  ("Address", "[INSERT ADDRESS]"),
                  ("Email and telephone", "[INSERT CONTACT DETAILS]"),
                  ("Signature", ""),
                  ("Date", ""),
              ],
              [2700, 6660])

    add_heading(doc, "Proposal checklist before submission", 2)
    for text in [
        "Replace every bracketed field and remove this draft instruction checklist.",
        "Reconcile the proposal against the final bidding documents, bid bulletins, approved budget, and required government forms.",
        "Confirm which unfinished features are mandatory base deliverables, optional items, or clarification matters.",
        "Complete the technical compliance matrix with truthful documentary support for every specification.",
        "Finalize the implementation schedule, biometric Bill of Quantities, commercial amounts, payment milestones, warranty, and SLA.",
        "Obtain legal, procurement, technical, data-privacy, security, and authorized-signatory review before submission.",
    ]:
        add_bullet(doc, text)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
